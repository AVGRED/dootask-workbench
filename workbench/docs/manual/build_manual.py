"""生成《星图 · 用户手册》DOCX。运行：python build_manual.py"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "微软雅黑"
MONO = "Consolas"
ACCENT = RGBColor(0x4F, 0x46, 0xE5)
GREY = RGBColor(0x6B, 0x72, 0x80)

document = Document()

# ---------- 页面与基础样式 ----------
section = document.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)


def set_cjk(style_or_run, size=None, bold=None, color=None):
    font = style_or_run.font
    font.name = FONT
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    element = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)


for style_name in ("List Bullet", "List Number"):
    try:
        set_cjk(document.styles[style_name], 10.5)
    except KeyError:
        pass

normal = document.styles["Normal"]
set_cjk(normal, 10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(4)

for level, size in ((1, 17), (2, 13.5), (3, 11.5)):
    style = document.styles[f"Heading {level}"]
    set_cjk(style, size, True, ACCENT if level == 1 else RGBColor(0x1F, 0x24, 0x33))
    style.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    style.paragraph_format.space_after = Pt(6)


# ---------- 工具函数 ----------
def para(text="", size=None, bold=False, color=None, align=None, space_after=None, indent=None):
    reset_numbering()
    p = document.add_paragraph()
    run = p.add_run(text)
    set_cjk(run, size, bold, color)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def bullet(text, level=0):
    reset_numbering()
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 + level * 0.6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_cjk(run)
    return p


_number_counter = {"value": 0}


def reset_numbering():
    _number_counter["value"] = 0


def numbered(text):
    """手动编号：每个列表都从 1 开始，避免 Word 的连续编号错乱。"""
    _number_counter["value"] += 1
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f"{_number_counter['value']}. {text}")
    set_cjk(run)
    return p


def code(text):
    reset_numbering()
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = MONO
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x33, 0x38, 0x45)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F4F6")
    p._p.get_or_add_pPr().append(shading)
    return p


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def set_fixed_layout(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for index, width in enumerate(widths):
        t.columns[index].width = Cm(width)
        for cell in t.columns[index].cells:
            cell.width = Cm(width)


def table(headers, rows, widths=None):
    reset_numbering()
    t = document.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = t.rows[0].cells
    for index, title in enumerate(headers):
        header[index].text = ""
        run = header[index].paragraphs[0].add_run(title)
        set_cjk(run, 9.5, True, RGBColor(0xFF, 0xFF, 0xFF))
        shade(header[index], "4F46E5")
        header[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            set_cjk(run, 9.5)
    if widths:
        set_fixed_layout(t, widths)
    # 表头跨页重复
    trPr = t.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    trPr.append(header_flag)
    # 单行内容不拆到两页
    for row in t.rows:
        row_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        row_pr.append(cant)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(title, text):
    reset_numbering()
    t = document.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    cell.text = ""
    p1 = cell.paragraphs[0]
    run = p1.add_run(title + "  ")
    set_cjk(run, 9.5, True, ACCENT)
    run2 = p1.add_run(text)
    set_cjk(run2, 9.5)
    shade(cell, "F5F6FF")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


_original_add_heading = document.add_heading


def _heading_with_reset(text, level=1):
    reset_numbering()
    return _original_add_heading(text, level)


document.add_heading = _heading_with_reset

# ---------- 页脚 ----------
footer_paragraph = section.footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_left = footer_paragraph.add_run("星图 · 用户手册　")
set_cjk(run_left, 9, False, GREY)
run_page = footer_paragraph.add_run()
set_cjk(run_page, 9, False, GREY)
fld_begin = OxmlElement("w:fldChar")
fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = "PAGE"
fld_end = OxmlElement("w:fldChar")
fld_end.set(qn("w:fldCharType"), "end")
run_page._r.append(fld_begin)
run_page._r.append(instr)
run_page._r.append(fld_end)

# ---------- 封面 ----------
for _ in range(3):
    document.add_paragraph()
logo = document.add_paragraph()
logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo.add_run().add_picture("../../brand/png/icon-192.png", width=Cm(3.2))
document.add_paragraph()
para("星图", 40, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para("WORKBENCH", 12, True, GREY, WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
para("用户手册", 26, True, RGBColor(0x1F, 0x24, 0x33), WD_ALIGN_PARAGRAPH.CENTER)
document.add_paragraph()
para("员工端 · 管理端 · 移动端 全功能使用说明", 12, False, GREY, WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(2):
    document.add_paragraph()
table(
    ["项目", "内容"],
    [
        ["适用系统", "https://yanyunxiaoguang.cn"],
        ["适用版本", "v2（星图）"],
        ["文档版本", "v1.0"],
        ["面向读者", "全体员工、项目负责人、管理员"],
    ],
    widths=[4, 11],
)
page_break()

# ---------- 目录 ----------
document.add_heading("目录", level=1)
toc = [
    ("第一章  系统简介", "星图是什么、三个端的分工、术语表"),
    ("第二章  开始使用", "访问地址、登录方式、界面结构、暗色模式"),
    ("第三章  员工端", "首页仪表盘、我的任务、项目、任务、看板、甘特、日报、星图、设置"),
    ("第四章  管理端", "状态墙、项目总览、每日日报、星图、通知、用户管理"),
    ("第五章  移动端", "手机访问、安装到桌面、功能范围"),
    ("第六章  权限说明", "员工与管理员的能做/不能做"),
    ("第七章  常见问题", "登录、星图、通知、日报等 12 个高频问题"),
    ("第八章  使用建议", "任务粒度、日报写法、状态维护等团队约定"),
]
for title, desc in toc:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title + "　")
    set_cjk(run, 11.5, True)
    run2 = p.add_run(desc)
    set_cjk(run2, 9.5, False, GREY)
page_break()

# ---------- 第一章 ----------
document.add_heading("第一章  系统简介", level=1)

document.add_heading("1.1 星图是什么", level=2)
para(
    "星图是一套面向小团队的项目与任务协作系统。它把项目、任务、里程碑、工时、日报、成员状态统一在一个工作台里，"
    "并用一张「星图」把人和人之间的协作关系可视化——每个人是一颗星，共同参与过项目的人之间会有连线。"
)
para("系统要解决的核心问题有三个：")
bullet("事情散了：项目在哪、谁的活、什么时候到期，集中看得到。")
bullet("进度黑箱：每天的状态、日报、任务进展都留痕，不用反复问「做到哪了」。")
bullet("协作靠猜：谁和谁在同一个项目上、跨了多少个项目，星图一眼看清。")

document.add_heading("1.2 三个端的分工", level=2)
table(
    ["端", "地址", "给谁用", "主要能力"],
    [
        ["员工端", "https://yanyunxiaoguang.cn/", "全体员工", "项目、任务、看板、甘特、日报、星图、个人待办、设置"],
        ["管理端", "https://yanyunxiaoguang.cn/admin/", "管理员 / 项目负责人", "状态墙、项目总览、每日日报、全公司星图、发布通知、用户管理"],
        ["移动端", "https://yanyunxiaoguang.cn/m/", "全体员工（手机）", "查看项目与任务、写日报、看通知、看星图"],
    ],
    widths=[2.6, 5.4, 3.6, 5.0],
)
callout("提示", "三个端共用同一套账号与数据，手机上看的东西和电脑上完全一致，不需要单独注册。")

page_break()
document.add_heading("1.3 术语表", level=2)
table(
    ["术语", "含义"],
    [
        ["项目", "一组相关任务的容器，有负责人、成员、起止时间和状态。"],
        ["任务", "可执行的最小工作项，有负责人、优先级、开始/截止时间、状态。"],
        ["里程碑", "项目中的关键节点，用来标记阶段性目标（如「1.0 发布」）。"],
        ["看板", "按状态分列的任务视图，拖拽卡片即可改状态。"],
        ["甘特图", "按时间轴展示任务排期与依赖关系，用来发现排期冲突。"],
        ["工时", "你在某个任务上投入的分钟数，用于容量与负载统计。"],
        ["日报", "每天的工作记录：今天做了什么、遇到什么问题、明天计划。"],
        ["状态", "员工当前的可用性：专注 / 会议 / 外出 / 请假 / 空闲。"],
        ["星图", "以人为节点的协作关系图，连线表示共同参与过项目。"],
        ["个人待办", "不挂在任何项目下的私人小事，只有自己能看到。"],
        ["通知", "系统或管理员发给你的站内消息，未读会显示红点。"],
    ],
    widths=[3.2, 12.0],
)
page_break()

# ---------- 第二章 ----------
document.add_heading("第二章  开始使用", level=1)

document.add_heading("2.1 访问地址", level=2)
table(
    ["用途", "地址"],
    [
        ["员工端（日常办公）", "https://yanyunxiaoguang.cn"],
        ["管理端（管理/查看全局）", "https://yanyunxiaoguang.cn/admin"],
        ["移动端（手机浏览器）", "https://yanyunxiaoguang.cn/m"],
    ],
    widths=[6.5, 8.0],
)
para("建议把员工端加入浏览器书签；手机上可以把页面「添加到主屏幕」，用起来接近一个独立 App。")

document.add_heading("2.2 登录", level=2)
para("打开员工端后会看到登录页，输入账号和密码即可。账号支持三种写法（任选其一）：")
numbered("工号：如 401004")
numbered("工号邮箱：如 401004@yanyunxiaoguang.cn")
numbered("昵称：如 毛腾辉")
callout("忘记密码", "系统不提供自助找回。请联系管理员在「用户管理」里帮你重置密码。")
para("勾选「记住账号」后，下次打开只需输入密码；公共电脑上请不要勾选。")

document.add_heading("2.3 第一次登录建议做的三件事", level=2)
numbered("改密码：右上角头像 → 设置 → 修改密码。新密码至少 8 位。")
numbered("传头像：设置 → 头像 → 上传新头像。头像会显示在状态墙和星图上。")
numbered("补资料：设置 → 基本信息，填写昵称和职位，方便同事认人。")

document.add_heading("2.4 界面结构", level=2)
para("登录后是员工端首页，从上到下分三块：")
table(
    ["区域", "包含内容", "说明"],
    [
        ["顶栏", "星图 Logo、导航、当前状态、主题切换、头像菜单", "导航包含：首页 / 项目 / 日报 / 星图 / 设置"],
        ["标题栏", "页面标题 + 页面级操作按钮", "如「新建项目」「写日报」"],
        ["内容区", "当前页面的主体内容", "首页是可自定义的仪表盘"],
    ],
    widths=[2.6, 6.6, 6.0],
)
para("右上角头像菜单里可以：进入设置、切换浅色/深色、退出登录。顶部中间的小圆点是你的状态（空闲/专注/会议…），点它可以快速改状态。")

document.add_heading("2.5 暗色模式", level=2)
para("顶栏最右侧的月亮/太阳图标可以切换浅色与深色主题。系统会记住你的选择；第一次访问时，会跟随操作系统的深色设置。")
page_break()

# ---------- 第三章 员工端 ----------
document.add_heading("第三章  员工端", level=1)

document.add_heading("3.1 首页仪表盘", level=2)
para("首页由若干「板块」组成，你可以按自己的工作习惯调整：")
bullet("拖动板块标题即可移动位置。")
bullet("拖动板块右下角可以调整大小。")
bullet("点「保存布局」后，布局会跟着你的账号走，换电脑打开也一样。")
bullet("点「恢复默认」回到系统初始布局。")
para("默认板块与其用途：")
table(
    ["板块", "内容", "怎么用"],
    [
        ["我的任务", "个人待办 + 项目任务", "顶部输入框写待办回车即存；下方是各项目里分给你的未完成任务"],
        ["今日日报", "今天是否已提交日报", "点「去写日报」直接跳转"],
        ["工时容量", "近 7 天已投入分钟数 / 可用分钟数", "用来判断自己是否被排满"],
        ["未读通知", "未读数量与最近几条", "点击进入通知列表"],
        ["项目概况", "我的项目数、未完成任务数、逾期数", "红色数字是逾期，需要优先处理"],
        ["最近动态", "团队最近的任务/评论等操作记录", "了解别人在做什么"],
        ["看板缩略", "每个项目的未完成任务数", "点进去直接到看板"],
        ["本周排期", "本周有开始或截止时间的任务", "提前发现撞车"],
        ["全局搜索", "搜项目 / 任务 / 人", "输入关键词回车"],
    ],
    widths=[3.0, 5.6, 6.6],
)

document.add_heading("3.2 我的任务", level=2)
para("「我的任务」板块里有两类东西，用途不一样：")
para("个人待办：")
bullet("不挂在任何项目下，相当于私人便签，别人看不到。")
bullet("在输入框写一句话，按回车即保存。")
bullet("勾选复选框标记完成，完成的条目会变灰并加删除线。")
bullet("鼠标移到条目右侧出现「×」，点击删除。")
bullet("适合记：临时事项、给自己排的今日清单、需要跟进但不属于项目的事。")
para("项目任务：")
bullet("来自各个项目分配给你的未完成任务。")
bullet("点击任务会跳到对应的项目任务页。")
bullet("右侧显示截止日期，逾期任务会在「项目概况」里变成红色计数。")

document.add_heading("3.3 项目", level=2)
para("点顶栏「项目」进入项目列表，可以看到自己参与的所有项目。")
para("新建项目：")
numbered("点右上角「新建项目」。")
numbered("填写项目名称、描述、起止时间。")
numbered("选择成员（可多选）。")
numbered("保存后自动跳转到项目详情。")
para("项目常见操作：")
table(
    ["操作", "位置", "说明"],
    [
        ["查看任务", "项目卡片 / 项目页", "进入该项目的任务列表"],
        ["看板", "项目页标签", "按状态分列查看与拖拽"],
        ["甘特图", "项目页标签", "按时间轴看排期与依赖"],
        ["邀请成员", "项目页成员区", "按昵称搜索添加"],
        ["退出项目", "项目页成员区", "非负责人可以退出，负责人不能退出"],
        ["删除项目", "项目页「更多」", "只有负责人或管理员能删；删除后任务一并隐藏"],
    ],
    widths=[3.0, 4.0, 8.2],
)

document.add_heading("3.4 任务", level=2)
para("任务是最常用的功能。新建任务时可以填：")
table(
    ["字段", "必填", "说明"],
    [
        ["任务名称", "是", "一句话说清要做什么，建议以动词开头"],
        ["负责人", "是", "默认是自己，可以改成别人"],
        ["优先级", "否", "低 / 中 / 高 / 紧急"],
        ["开始 / 截止时间", "否", "甘特图与逾期判断依赖这两个字段"],
        ["描述", "否", "补充背景、验收标准、参考链接"],
        ["标签", "否", "用于分类过滤"],
        ["里程碑", "否", "把任务挂到项目的某个节点下"],
    ],
    widths=[3.2, 1.6, 10.4],
)
para("任务详情页还可以：")
bullet("添加子任务：把大任务拆成可勾选的小步骤。")
bullet("设置依赖：前置任务没完成时，甘特图会提示冲突。")
bullet("写评论：支持 @ 同事，被 @ 的人会收到通知。")
bullet("上传附件：单个文件有大小限制，超出会提示。")
bullet("记录工时：在任务上登记投入的分钟数，用于统计。")
bullet("修改状态：待开始 / 进行中 / 已完成 / 阻塞。")
callout("注意", "任务完成后不会立即消失，会进入「已完成」筛选里；逾期未完成的任务会一直在你的逾期统计中。")

document.add_heading("3.5 看板", level=2)
para("看板把任务按状态分成几列，拖动卡片就能改状态。适合每天站会时快速过一遍。")
bullet("拖动卡片到目标列，状态自动更新。")
bullet("卡片上显示负责人、优先级、截止时间。")
bullet("列头可以看到每列有多少任务。")

document.add_heading("3.6 甘特图", level=2)
para("甘特图按时间轴展示任务的排期，用来回答「这周排得下吗」。")
bullet("横条长度 = 任务的开始到截止。")
bullet("任务之间的依赖会用连线表示。")
bullet("同一人被安排了重叠的任务时，容易看出资源冲突。")

document.add_heading("3.7 日报", level=2)
para("点顶栏「日报」进入日报页。每天写一篇，内容通常包括：")
numbered("今天做了什么（对应到具体任务更好）。")
numbered("遇到的问题 / 需要协调的事。")
numbered("明天计划做什么。")
para("写完后点「提交」。提交后当天会显示「已提交」；管理员可以在「每日日报」里看到谁还没交。")
callout("补交", "错过当天可以补写，日报按日期归档，补写时注意选对日期。")

document.add_heading("3.8 星图", level=2)
para("点顶栏「星图」，进入你自己的协作关系图：")
bullet("每颗星代表一个和你合作过的同事（包括你自己）。")
bullet("两颗星之间的连线表示你们共同参与过项目。")
bullet("连线上的文字标出合作的项目名称。")
bullet("星越大，表示参与的项目越多。")
bullet("鼠标悬停某颗星，会高亮与它相关的连线；悬停连线可以看到合作项目。")
bullet("节点里显示对方头像（没上传头像时显示默认剪影），下方是名字。")
bullet("你只能看到与自己有合作的人，看不到全公司。")
para("管理端的星图是另一种视角：能看到全公司所有人，以及所有协作关系。")

document.add_heading("3.9 设置", level=2)
para("点顶栏「设置」（或头像菜单）进入个人设置，包含三块：")
table(
    ["卡片", "能改什么", "说明"],
    [
        ["基本信息", "昵称、邮箱、职位", "邮箱要符合格式且不能和别人重复；改完立即生效"],
        ["修改密码", "原密码、新密码、确认", "新密码至少 8 位；改完下次用新密码登录"],
        ["头像", "上传新头像", "支持 JPEG / PNG，建议正方形、小于 2MB"],
    ],
    widths=[3.0, 5.0, 7.2],
)

document.add_heading("3.10 通知", level=2)
para("有未读消息时，顶栏会显示红点，首页「未读通知」板块也能看到数量。")
bullet("任务被你负责、评论里 @ 你、管理员发通知，都会产生通知。")
bullet("点开后可以逐条标记已读，也可以一键全部已读。")
bullet("通知不会重复提醒，已读后红点消失。")
page_break()

# ---------- 第四章 管理端 ----------
document.add_heading("第四章  管理端", level=1)
para("管理端地址：https://yanyunxiaoguang.cn/admin。只有角色为「管理员」的账号能登录。")
para("顶栏有五个标签页：状态墙、项目总览、每日日报、星图、通知；右上角另有用户管理与设置。")

document.add_heading("4.1 状态墙", level=2)
para("状态墙把每个人的当前状态做成卡片，适合快速看「现在谁有空」。")
bullet("顶部可按状态筛选：全部 / 专注 / 会议 / 外出 / 请假 / 空闲。")
bullet("右上角搜索框支持按姓名、ID、任务筛选。")
bullet("每张卡片显示：头像与昵称、当前状态、状态备注、正在做的任务、今日日报是否提交、近 7 天工时。")
bullet("状态墙只展示员工，管理员账号不会出现在状态墙上。")
callout("说明", "员工的状态由员工自己在顶栏切换；管理员在状态墙看到的是实时结果。")

document.add_heading("4.2 项目总览", level=2)
para("项目总览列出全公司的项目，方便横向对比进度。")
bullet("每个项目显示负责人、成员数、任务总数、已完成数、逾期数、完成百分比。")
bullet("可以按状态、负责人筛选，也可以搜索项目名。")
bullet("点项目可以查看详情与其任务。")

document.add_heading("4.3 每日日报", level=2)
para("用来检查谁没交日报。")
bullet("选择日期后，左边是已提交的日报列表，右边是未提交人员名单。")
bullet("点开一份日报可以看到完整内容。")
bullet("名单可以复制，方便在群里提醒。")

document.add_heading("4.4 星图（全公司）", level=2)
para("管理端的星图能看到全部员工与全部协作关系：")
bullet("节点 = 员工，节点大小 = 参与项目数。")
bullet("连线 = 共同参与过项目，线上的文字标注项目名称。")
bullet("鼠标悬停可高亮相关连线并显示合作详情。")
bullet("人数或关系过多时，图会做截断并给出提示。")

document.add_heading("4.5 通知（发布公告）", level=2)
para("管理员可以向全体员工发布通知：")
numbered("进入「通知」标签页。")
numbered("填写标题（最多 120 字）与内容（最多 1000 字）。")
numbered("点「发布给全部员工」。")
numbered("下方「历史通知」会保留发布记录，含发布人与时间。")
callout("效果", "发布后每位员工的站内通知里会收到一条消息，未读会显示红点。当前系统没有邮件/短信通道，只在站内提醒。")

document.add_heading("4.6 用户管理", level=2)
para("右上角「用户管理」用于维护账号：")
table(
    ["操作", "说明"],
    [
        ["新建用户", "填写邮箱、昵称、初始密码；邮箱不能重复"],
        ["搜索", "按邮箱或昵称搜索"],
        ["改角色", "员工 ↔ 管理员，改完立即生效"],
        ["重置密码", "管理员给用户设置新密码，用户登录后可自行修改"],
        ["删除用户", "单条删除或勾选后批量删除；系统禁止删除最后一个管理员"],
    ],
    widths=[3.2, 11.8],
)
callout("注意", "删除用户是软删除：账号无法登录，但历史数据保留，便于追溯。")

document.add_heading("4.7 设置（管理端）", level=2)
para("管理端的「设置」与员工端相同，用于修改自己的昵称、邮箱、职位、密码和头像。")
page_break()

# ---------- 第五章 移动端 ----------
document.add_heading("第五章  移动端", level=1)

document.add_heading("5.1 访问", level=2)
para("手机浏览器打开 https://yanyunxiaoguang.cn/m，用与电脑端相同的账号登录。")

document.add_heading("5.2 安装到桌面", level=2)
para("移动端支持「添加到主屏幕」，安装后体验接近原生 App：")
bullet("iPhone：Safari → 分享 → 添加到主屏幕。")
bullet("Android：Chrome → 右上角菜单 → 添加到主屏幕 / 安装应用。")
bullet("安装后图标为星图 Logo，打开后没有浏览器地址栏。")

document.add_heading("5.3 能做什么", level=2)
table(
    ["功能", "支持情况"],
    [
        ["查看项目与任务", "支持"],
        ["更新任务状态", "支持"],
        ["写日报", "支持"],
        ["查看通知", "支持"],
        ["查看星图", "支持"],
        ["个人待办", "支持"],
        ["看板拖拽", "不支持（请在电脑端操作）"],
        ["甘特图", "仅查看"],
    ],
    widths=[6.0, 9.0],
)
page_break()

# ---------- 第六章 权限 ----------
document.add_heading("第六章  权限说明", level=1)
para("系统只有两种角色：员工（staff）与管理员（admin）。")
table(
    ["能力", "员工", "管理员"],
    [
        ["查看自己参与的项目与任务", "可以", "可以"],
        ["创建项目、任务、日报、待办", "可以", "可以"],
        ["查看他人的项目与任务", "仅限共同项目", "全部"],
        ["查看全公司状态墙", "不可以", "可以"],
        ["查看全公司日报提交情况", "不可以", "可以"],
        ["查看全公司星图", "不可以（只看自己的协作圈）", "可以"],
        ["发布全员通知", "不可以", "可以"],
        ["用户管理（新建/删除/改角色/改密码）", "不可以", "可以"],
        ["修改自己的资料与密码", "可以", "可以"],
    ],
    widths=[7.0, 4.2, 3.8],
)
callout("重要", "任务、评论、附件等资源的访问都按「是否属于该项目成员」校验：非成员访问会返回「任务不存在」，不会泄露内容。")

document.add_heading("6.1 管理员注意事项", level=2)
bullet("管理员账号不会出现在状态墙上，避免与员工混淆。")
bullet("多个管理员权限相同，没有「超级管理员」之分。")
bullet("删除用户前请确认其项目已有其他负责人。")
bullet("发布通知会发给全体在职员工，内容请先确认措辞。")
page_break()

# ---------- 第七章 FAQ ----------
document.add_heading("第七章  常见问题", level=1)
faqs = [
    (
        "忘记密码怎么办？",
        "系统不支持自助找回。请找管理员在「用户管理 → 重置密码」里设置新密码，登录后自己再改。",
    ),
    (
        "用工号登录提示「账号或密码错误」？",
        "先确认密码是否被管理员重置过；再确认工号是否正确（如 401004）。工号登录依赖邮箱前缀，"
        "如果管理员改过你的邮箱，请改用新邮箱或昵称登录。",
    ),
    (
        "星图上只有自己一颗星，没有连线？",
        "星图的连线来自「共同参与的项目」。当你和别人在同一个项目里时才会连线。让负责人把你加入项目即可。",
    ),
    (
        "星图上看不到某个同事？",
        "员工端只显示与你合作过的人。如果你们从未在同一个项目里，就不会出现。",
    ),
    (
        "状态墙上没有某个管理员的卡片？",
        "这是设计如此：状态墙只展示员工，管理员不上墙。",
    ),
    (
        "我改了昵称/头像，为什么别处还是旧的？",
        "部分页面有缓存，刷新页面即可；状态墙与星图的头像是实时读取的。",
    ),
    (
        "日报可以补交吗？",
        "可以。日报按日期归档，打开日报页选择对应日期即可补写。",
    ),
    (
        "为什么我的任务没有出现在「我的任务」里？",
        "「我的任务」只显示分配给你、且尚未完成的任务。如果任务已完成、或被移出项目，就不会显示。",
    ),
    (
        "附件上传失败？",
        "单个文件有大小上限（约 20MB）。请压缩图片或分包上传。",
    ),
    (
        "手机收不到通知提醒？",
        "当前只有站内通知，没有系统级推送。把移动端添加到主屏幕后，打开时能看到未读红点。",
    ),
    (
        "能删除项目吗？",
        "只有项目负责人或管理员能删除。删除后项目与任务对其他成员隐藏，数据保留在库中。",
    ),
    (
        "换电脑后我的自定义首页布局还在吗？",
        "在。布局保存在你的账号下，与设备无关；但浏览器主题（浅色/深色）保存在本地，换电脑需要重新选择。",
    ),
]
for index, (question, answer) in enumerate(faqs, start=1):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"Q{index}. {question}")
    set_cjk(run, 10.5, True, RGBColor(0x1F, 0x24, 0x33))
    para(answer, indent=0.4)
page_break()

# ---------- 第八章 建议 ----------
document.add_heading("第八章  使用建议", level=1)
para("工具用得好不好，更多取决于团队约定。以下是推荐做法。")

document.add_heading("8.1 任务粒度", level=2)
bullet("一个任务最好在 0.5～2 天能做完；超过 3 天请拆成子任务。")
bullet("任务名称用动词开头：「完成登录页改版」比「登录页」更好。")
bullet("验收标准写进描述，避免做完才发现理解不一致。")

document.add_heading("8.2 每天的三分钟", level=2)
numbered("早上：打开首页，看「本周排期」和「我的任务」，把今天要做的勾出来。")
numbered("白天：开始一件事就切换状态（专注/会议），让同事知道你在忙什么。")
numbered("下班前：写日报，勾掉完成的任务，把没做完的挪到明天。")

document.add_heading("8.3 日报写法", level=2)
para("好日报的标准是「别人不看聊天记录也能看懂」：")
bullet("今天做了什么：写结果，不写过程（「接口联调完成」而不是「一直在联调」）。")
bullet("遇到问题：写清楚卡在哪、需要谁配合。")
bullet("明天计划：写 1～3 件最重要的事。")

document.add_heading("8.4 状态维护", level=2)
table(
    ["状态", "什么时候用"],
    [
        ["专注", "正在写代码/写文档，不希望被打断"],
        ["会议", "在开会，无法及时回复"],
        ["外出", "不在工位"],
        ["请假", "当天不在岗"],
        ["空闲", "可以接新任务、可以找人"],
    ],
    widths=[3.0, 12.0],
)

document.add_heading("8.5 项目负责人要做的三件事", level=2)
numbered("每周一检查项目总览：有没有逾期任务、有没有人被排满。")
numbered("任务拆得够细，负责人明确，截止时间真实。")
numbered("项目结束时清理未完成任务（关闭或转交），不要留僵尸任务。")

document.add_heading("附录：技术支持", level=2)
para("遇到系统问题时，请提供以下信息，便于快速定位：")
bullet("发生时间（精确到分钟）")
bullet("操作路径（在哪个页面点了什么）")
bullet("报错截图或提示文字")
bullet("账号（工号）与使用的端（电脑浏览器 / 手机）")
para("管理员联系方式请以团队内部通知为准。")

document.save("星图-用户手册.docx")
print("saved 星图-用户手册.docx")
